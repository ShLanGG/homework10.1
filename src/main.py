"""
Главный модуль программы, реализующий интерфейс командной строки
для фильтрации и анализа банковских транзакций.
"""
import json
import csv
from typing import List, Optional

from src.search import search_transactions_by_description

# Подставьте импорт существующих загрузчиков из соответствующего модуля.
# Например:
# from src.file_reader import load_transactions_json, load_transactions_csv, load_transactions_xlsx
# Ниже заготовки на случай отсутствия готовых модулей.
def load_transactions_json(path: str) -> List[dict]:
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)

def load_transactions_csv(path: str) -> List[dict]:
    with open(path, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        return list(reader)

def load_transactions_xlsx(path: str) -> List[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    transactions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        transactions.append(dict(zip(headers, row)))
    return transactions


def filter_by_status(transactions: List[dict], status: str) -> List[dict]:
    """Фильтрует список транзакций по статусу (без учёта регистра)."""
    return [t for t in transactions if t.get('state', '').upper() == status.upper()]


def input_yes_no(prompt: str) -> bool:
    """Запрашивает у пользователя Да/Нет и возвращает True/False."""
    while True:
        answer = input(prompt).strip().lower()
        if answer in ('да', 'yes', 'y'):
            return True
        if answer in ('нет', 'no', 'n'):
            return False
        print('Пожалуйста, введите "Да" или "Нет".')


def main():
    """
    Основная логика программы.
    Предоставляет пользователю интерфейс для выбора файла, фильтрации,
    сортировки и вывода банковских транзакций.
    """
    print("Привет! Добро пожаловать в программу работы с банковскими транзакциями.")
    print("Выберите необходимый пункт меню:")
    print("1. Получить информацию о транзакциях из JSON-файла")
    print("2. Получить информацию о транзакциях из CSV-файла")
    print("3. Получить информацию о транзакциях из XLSX-файла")
    
    file_type = input("Введите номер: ").strip()
    
    loaders = {
        '1': ('JSON-файл.', load_transactions_json),
        '2': ('CSV-файл.', load_transactions_csv),
        '3': ('XLSX-файл.', load_transactions_xlsx)
    }
    
    if file_type not in loaders:
        print("Некорректный выбор. Программа завершена.")
        return
    
    file_prompt, loader = loaders[file_type]
    print(f"\nДля обработки выбран {file_prompt}")
    file_path = input("Введите путь к файлу: ").strip()
    
    try:
        transactions = loader(file_path)
    except Exception as e:
        print(f"Ошибка при чтении файла: {e}")
        return
    
    # ---------- Фильтрация по статусу ----------
    valid_statuses = {'EXECUTED', 'CANCELED', 'PENDING'}
    while True:
        status_input = input(
            "\nВведите статус, по которому необходимо выполнить фильтрацию.\n"
            "Доступные для фильтровки статусы: EXECUTED, CANCELED, PENDING\n>> "
        ).strip()
        if status_input.upper() in valid_statuses:
            break
        print(f'Статус операции "{status_input}" недоступен.')
    
    transactions = filter_by_status(transactions, status_input.upper())
    print(f'\nОперации отфильтрованы по статусу "{status_input.upper()}"\n')
    
    # ---------- Сортировка по дате ----------
    if input_yes_no("Отсортировать операции по дате? Да/Нет\n>> "):
        order = input("Отсортировать по возрастанию или по убыванию?\n>> ").strip().lower()
        reverse = 'убыванию' in order or 'убыв' in order
        # Для сортировки используем поле 'date' в формате ДД.ММ.ГГГГ,
        # переведём в объекты datetime для правильного сравнения
        from datetime import datetime
        try:
            transactions.sort(
                key=lambda t: datetime.strptime(t['date'], '%d.%m.%Y'),
                reverse=reverse
            )
        except (KeyError, ValueError):
            print("Не удалось отсортировать: неверный формат даты.")
    
    # ---------- Только рублёвые ----------
    if input_yes_no("\nВыводить только рублевые транзакции? Да/Нет\n>> "):
        transactions = [
            t for t in transactions
            if 'operationAmount' in t
            and t['operationAmount'].get('currency', {}).get('code', '').upper() == 'RUB'
        ]
    
    # ---------- Фильтрация по слову в описании ----------
    if input_yes_no("\nОтфильтровать список транзакций по определенному слову в описании? Да/Нет\n>> "):
        search_word = input("Введите слово для поиска: ").strip()
        transactions = search_transactions_by_description(transactions, search_word)
    
    # ---------- Вывод результата ----------
    print("\nРаспечатываю итоговый список транзакций...\n")
    if not transactions:
        print("Не найдено ни одной транзакции, подходящей под ваши условия фильтрации")
        return
    
    print(f"Всего банковских операций в выборке: {len(transactions)}\n")
    for t in transactions:
        date = t.get('date', '—')
        description = t.get('description', '—')
        from_ = t.get('from', '')
        to_ = t.get('to', '')
        amount = t.get('operationAmount', {}).get('amount', '')
        currency = t.get('operationAmount', {}).get('currency', {}).get('code', '')
        
        if from_ and to_:
            direction = f"{from_} -> {to_}"
        elif from_:
            direction = f"{from_} -> —"
        elif to_:
            direction = f"— -> {to_}"
        else:
            direction = ""
        
        print(f"{date} {description}")
        if direction:
            print(direction)
        print(f"Сумма: {amount} {currency}")
        print()


if __name__ == "__main__":
    main()